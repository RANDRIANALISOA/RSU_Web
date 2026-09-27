# -*- coding: utf-8 -*-
"""
maj_zones_31districts.py — Mise à jour ciblée du référentiel `zones` (région,
district, commune, fokontany) pour les 31 districts de la vague RSU, à partir du
fichier « Base compilée » (MAJ_ZONES_31DISTRICTS.xlsx).

POURQUOI un script séparé de zones.py :
    `zones.charger_excel_vers_db` recharge TOUT le pays depuis FKT_ampiasan_SS
    (20 256 fokontany). Ici on ne touche QUE les 31 districts présents dans le
    fichier fourni ; les 89 autres districts restent intacts.

SÉMANTIQUE = REMPLACEMENT (le fichier fait foi, décision utilisateur 2026-09-26) :
    pour chaque district du fichier, on SUPPRIME ses communes/fokontany actuels puis
    on réinsère EXACTEMENT ceux du fichier. Les noms de district et de région sont
    mis à jour d'après le fichier. On charge TOUT tel quel, y compris les codes
    fokontany qui ne s'emboîtent pas dans leur commune (cf. rapport d'anomalies :
    543 « même district, commune différente » + 52 « district différent »). La
    hiérarchie retenue est celle des COLONNES du fichier (commune de la ligne,
    district de la ligne), pas celle déduite des codes.

Le fichier n'a pas de code région/province : ils sont déduits des codes
(région = code_district // 100 ; province = code_région // 10), conforme au schéma
du référentiel (province 1 chiffre, région 2, district 4, commune 6, fokontany 8).
`commune."nombreMenage"` (projection ménages) est PRÉSERVÉ par code de commune.

⚠️ Un rechargement complet ultérieur (`python zones.py`) écraserait ces 31 districts
   par FKT_ampiasan_SS : relancer ce script après, ou mettre à jour le fichier maître.
⚠️ Les clés étrangères de `zones` ne sont PAS appliquées par SQLite : d'anciennes
   références (den_menage, superviseur_commune…) vers un fokontany/commune supprimé
   deviennent orphelines sans erreur — attendu, « le fichier fait foi ».

Usage :
    python maj_zones_31districts.py            # applique la mise à jour
    python maj_zones_31districts.py --dry-run  # simule et affiche le bilan
    python maj_zones_31districts.py <fichier.xlsx>
"""
from __future__ import annotations

import os
import sys

import db_source
import zones

FEUILLE = "Base compilée"
XLSX_DEFAUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "MAJ_ZONES_31DISTRICTS.xlsx")


def _int(v):
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return None


def lire_fichier(xlsx_path: str):
    """Lit « Base compilée ». Renvoie (regions, districts, communes, fokontany) :
      regions   : {code_region: nom}
      districts : {code_district: (nom, code_region)}
      communes  : {code_commune: (nom, code_district)}   (1re occurrence gardée)
      fokontany : [(code_fokontany, nom, code_commune), ...] (codes uniques)
    La hiérarchie vient des COLONNES (commune/district de chaque ligne)."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[FEUILLE] if FEUILLE in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    next(it)                                   # en-tête
    regions, districts, communes = {}, {}, {}
    fokontany, vus_fkt = [], set()
    for r in it:
        if not r or _int(r[1]) is None:
            continue
        reg_nom = (str(r[0]).strip() if r[0] is not None else "")
        cd, ld = _int(r[1]), (str(r[2]).strip() if r[2] is not None else "")
        cc, lc = _int(r[3]), (str(r[4]).strip() if r[4] is not None else "")
        cf, lf = _int(r[5]), (str(r[6]).strip() if r[6] is not None else "")
        if cd is None or cc is None or cf is None:
            continue
        rc = cd // 100                          # code région = 2 premiers chiffres
        regions[rc] = reg_nom
        districts[cd] = (ld, rc)
        communes.setdefault(cc, (lc, cd))
        if cf not in vus_fkt:
            vus_fkt.add(cf)
            fokontany.append((cf, lf, cc))
    return regions, districts, communes, fokontany


def appliquer(conn, xlsx_path: str = XLSX_DEFAUT, dry_run: bool = False, log=print):
    """Applique la mise à jour (remplacement des 31 districts). Renvoie un bilan."""
    zones.assurer_colonne_menages(conn)         # garantit commune."nombreMenage"
    regions, districts, communes, fokontany = lire_fichier(xlsx_path)
    d31 = sorted(districts)
    log(f"Fichier : {len(regions)} régions, {len(d31)} districts, "
        f"{len(communes)} communes, {len(fokontany)} fokontany.")

    ph = db_source._placeholder(conn)
    cur = conn.cursor()

    # Sauvegarde du nombre de ménages par commune (réappliqué par code après réinsertion).
    cur.execute('SELECT code_commune, "nombreMenage" FROM "commune"')
    menages = {c: m for c, m in cur.fetchall() if m is not None}

    # Décompte AVANT (sur les 31 districts) pour le bilan.
    marks = ",".join([ph] * len(d31))
    cur.execute(f'SELECT COUNT(*) FROM "commune" WHERE code_district IN ({marks})', d31)
    com_avant = cur.fetchone()[0]
    cur.execute(f'SELECT COUNT(*) FROM "fokontany" f JOIN "commune" c '
                f'ON f.code_commune = c.code_commune '
                f'WHERE c.code_district IN ({marks})', d31)
    fkt_avant = cur.fetchone()[0]

    if dry_run:
        log(f"[DRY-RUN] Districts visés : {d31}")
        log(f"[DRY-RUN] Communes  : {com_avant} en base -> {len(communes)} (fichier)")
        log(f"[DRY-RUN] Fokontany : {fkt_avant} en base -> {len(fokontany)} (fichier)")
        return {"dry_run": True, "districts": len(d31),
                "communes_avant": com_avant, "communes_apres": len(communes),
                "fokontany_avant": fkt_avant, "fokontany_apres": len(fokontany)}

    try:
        # 1) Régions (+ provinces déduites) : upsert du nom.
        for rc, nom in sorted(regions.items()):
            pc = rc // 10
            cur.execute(f'SELECT 1 FROM "province" WHERE code_province={ph}', (pc,))
            if cur.fetchone() is None:
                cur.execute(f'INSERT INTO "province" (code_province, nom) '
                            f'VALUES ({ph},{ph})', (pc, f"Province {pc}"))
            cur.execute(f'SELECT 1 FROM "region" WHERE code_region={ph}', (rc,))
            if cur.fetchone() is None:
                cur.execute(f'INSERT INTO "region" (code_region, nom, code_province) '
                            f'VALUES ({ph},{ph},{ph})', (rc, nom, pc))
            else:
                cur.execute(f'UPDATE "region" SET nom={ph}, code_province={ph} '
                            f'WHERE code_region={ph}', (nom, pc, rc))

        # 2) Districts : upsert du nom + rattachement région ; puis PURGE de leurs
        #    communes/fokontany (on supprime d'abord TOUS les 31 avant de réinsérer,
        #    pour tolérer un code commune qui « déménage » d'un district à l'autre).
        for cd in d31:
            ld, rc = districts[cd]
            cur.execute(f'SELECT 1 FROM "district" WHERE code_district={ph}', (cd,))
            if cur.fetchone() is None:
                cur.execute(f'INSERT INTO "district" (code_district, nom, code_region) '
                            f'VALUES ({ph},{ph},{ph})', (cd, ld, rc))
            else:
                cur.execute(f'UPDATE "district" SET nom={ph}, code_region={ph} '
                            f'WHERE code_district={ph}', (ld, rc, cd))
            cur.execute(f'DELETE FROM "fokontany" WHERE code_commune IN '
                        f'(SELECT code_commune FROM "commune" WHERE code_district={ph})',
                        (cd,))
            cur.execute(f'DELETE FROM "commune" WHERE code_district={ph}', (cd,))

        # 3) Communes du fichier (nombreMenage réappliqué par code si connu).
        cur.executemany(
            f'INSERT INTO "commune" (code_commune, nom, code_district, "nombreMenage") '
            f'VALUES ({ph},{ph},{ph},{ph})',
            [(cc, lc, cd, menages.get(cc)) for cc, (lc, cd) in communes.items()])

        # 4) Fokontany du fichier.
        cur.executemany(
            f'INSERT INTO "fokontany" (code_fokontany, nom, code_commune) '
            f'VALUES ({ph},{ph},{ph})', fokontany)

        conn.commit()
    except Exception:
        conn.rollback()
        raise

    log(f"OK : {len(d31)} districts remplacés. "
        f"Communes {com_avant} -> {len(communes)} ; "
        f"Fokontany {fkt_avant} -> {len(fokontany)}.")
    return {"districts": len(d31),
            "communes_avant": com_avant, "communes_apres": len(communes),
            "fokontany_avant": fkt_avant, "fokontany_apres": len(fokontany)}


def main() -> int:
    args = [a for a in sys.argv[1:]]
    dry = "--dry-run" in args
    args = [a for a in args if a != "--dry-run"]
    xlsx = args[0] if args else XLSX_DEFAUT
    if not os.path.exists(xlsx):
        print(f"Fichier introuvable : {xlsx}")
        return 1
    conn = db_source.connect()
    try:
        appliquer(conn, xlsx, dry_run=dry)
    finally:
        conn.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
