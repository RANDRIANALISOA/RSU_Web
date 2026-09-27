# -*- coding: utf-8 -*-
"""
anomalies_zones.py — Détecte, dans « Base compilée » (MAJ_ZONES_31DISTRICTS.xlsx),
les fokontany dont le code ne s'emboîte PAS dans le code de leur commune, et écrit un
classeur `anomalies_zones_31districts.xlsx` (3 feuilles : Résumé + 2 catégories).

Règle : un code fokontany « normal » commence par le code de sa commune
(commune 110101 -> fokontany 11010101…). Deux écarts :
  • « même district, commune différente » (renumérotation) ;
  • « district différent » (code pointant vers un autre district — contamination probable).

Usage :
    python anomalies_zones.py [fichier.xlsx] [sortie.xlsx]
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FEUILLE = "Base compilée"
XLSX_DEFAUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "MAJ_ZONES_31DISTRICTS.xlsx")
SORTIE_DEFAUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             "anomalies_zones_31districts.xlsx")
HDR = ["Région", "Code District", "Libelle District", "Code Commune",
       "Libelle Commune", "Code Fokontany", "Libelle Fokontany", "Long. code fkt",
       "Commune impliquée par le code", "District impliqué par le code"]


def _implique(cf, cc):
    """(commune, district) impliqués par le code fokontany (8 ou 9 chiffres)."""
    if len(str(cf)) == 9:
        return cf // 1000, cf // 100000
    return cf // 100, cf // 10000


def detecter(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[FEUILLE] if FEUILLE in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    next(it)
    total, diff_com, diff_dist = 0, [], []
    for r in it:
        if not r or r[1] is None:
            continue
        total += 1
        reg, cd, ld, cc, lc, cf, lf = r[:7]
        com_i, dist_i = _implique(int(cf), int(cc))
        if com_i == int(cc):
            continue
        ligne = [reg, cd, ld, cc, lc, cf, lf, len(str(cf)), com_i, dist_i]
        (diff_com if dist_i == int(cd) else diff_dist).append(ligne)
    return total, diff_com, diff_dist


def ecrire(sortie, total, diff_com, diff_dist):
    titre = Font(bold=True, color="FFFFFFFF")
    fill = PatternFill("solid", fgColor="FF1F4E79")
    th = Font(bold=True, color="FF1F2937")
    thf = PatternFill("solid", fgColor="FFDCE6F1")
    bd = Border(*[Side(style="thin", color="FFBFBFBF")] * 4)
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "Résumé"
    ws0["A1"] = "Anomalies du fichier « Base compilée » (31 districts)"
    ws0["A1"].font = Font(bold=True, size=13)
    ws0.append([])
    ws0.append(["Indicateur", "Nombre"])
    for c in ws0[3]:
        c.font, c.fill, c.border = th, thf, bd
    for a, b in (
            ("Total lignes (fokontany)", total),
            ("Fokontany correctement emboîtés", total - len(diff_com) - len(diff_dist)),
            ("ANOMALIE 1 — même district, code commune différent", len(diff_com)),
            ("ANOMALIE 2 — code pointant vers un AUTRE district (grave)", len(diff_dist))):
        ws0.append([a, b])
        for c in ws0[ws0.max_row]:
            c.border = bd
    ws0.column_dimensions["A"].width = 62
    ws0.column_dimensions["B"].width = 12

    def feuille(nom, data, note):
        w = wb.create_sheet(nom)
        w.append([note])
        w["A1"].font = Font(italic=True, color="FFB45309")
        w.append([])
        w.append(HDR)
        for c in w[3]:
            c.font, c.fill = titre, fill
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            c.border = bd
        for ligne in sorted(data, key=lambda x: (x[1], x[3], x[5])):
            w.append(ligne)
            for c in w[w.max_row]:
                c.border = bd
        for i, wd in enumerate([16, 12, 20, 13, 22, 15, 30, 10, 16, 16], 1):
            w.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wd
        w.freeze_panes = "A4"

    feuille("2_district_different", diff_dist,
            "ANOMALIE GRAVE : le code du fokontany appartient à un AUTRE district que "
            "celui indiqué (contamination probable). À corriger à la source.")
    feuille("1_commune_differente", diff_com,
            "Le fokontany reste dans le bon district, mais son code ne commence pas par "
            "le code de la commune indiquée (renumérotation / commune différente).")
    wb.save(sortie)


def main():
    args = sys.argv[1:]
    xlsx = args[0] if len(args) >= 1 else XLSX_DEFAUT
    sortie = args[1] if len(args) >= 2 else SORTIE_DEFAUT
    total, diff_com, diff_dist = detecter(xlsx)
    ecrire(sortie, total, diff_com, diff_dist)
    print(f"{total} fokontany | même district : {len(diff_com)} | "
          f"district différent : {len(diff_dist)}")
    print(f"Écrit : {sortie}")


if __name__ == "__main__":
    main()
